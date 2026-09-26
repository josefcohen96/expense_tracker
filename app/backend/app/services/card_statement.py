"""Import of the monthly Max credit-card statement (transaction-details_export_*.xlsx).

Nothing here is stored: the payer is guessed from the holder's name above the
header row, charges the app already has are recognised by date and amount (or,
for rows a recurrence materialised, by amount within a few days), and each
merchant's category is learned from the newest transaction whose notes equal
the merchant name.
"""
from __future__ import annotations

import re
import sqlite3
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any, Iterable, Optional

from openpyxl import load_workbook

from .access import display_name
from .people import find_person, household

# Max's `קטגוריה` column → app category names, first existing name wins.
MAX_CATEGORY_MAP: dict[str, tuple[str, ...]] = {
    "מסעדות, קפה וברים": ("אוכל בחוץ",),
    "מזון וצריכה": ("הוצאות בית",),
    "דלק, חשמל וגז": ("רכב",),
    "תחבורה ורכבים": ("תחבורה",),
    "שירותי תקשורת": ("הוצאות בית",),
    "חשמל ומחשבים": ("הוצאות בית",),
    "עירייה וממשלה": ("הוצאות בית",),
    "ביטוח": ("הוצאות בית",),
    "פנאי, בידור וספורט": ("פנאי",),
    "אופנה": ("פנאי",),
    "תיירות ונופש": ("פנאי",),
    "טיסות": ("פנאי",),
    "בריאות": ("בריאות",),
    "רפואה": ("בריאות",),
    "בתי מרקחת": ("בריאות",),
}

FALLBACK_CATEGORY = "הוצאות בית"

# Days either side of a file date within which a recurrence-booked row still counts.
RECURRING_WINDOW_DAYS = 3

HEADER_DATE = "תאריך עסקה"
HEADER_AMOUNT = "סכום חיוב"
HEADER_COLUMNS = {
    "merchant": "שם בית העסק",
    "max_category": "קטגוריה",
    "card_last4": "4 ספרות אחרונות של כרטיס האשראי",
    "currency": "מטבע חיוב",
    "charge_date": "תאריך חיוב",
    "remarks": "הערות",
}

_MONTH_RE = re.compile(r"^\d{2}/\d{4}$")
_DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d")


def normalise_merchant(value: Any) -> str:
    """Whitespace collapsed to single spaces and stripped."""
    return " ".join(str(value or "").split())


def _cell_text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _parse_date(value: Any) -> Optional[str]:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _cell_text(value)
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _parse_amount(value: Any) -> Optional[float]:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    text = _cell_text(value).replace("₪", "").replace(",", "").strip()
    if not text:
        return None
    try:
        return round(float(text), 2)
    except ValueError:
        return None


def _find_header(rows: list[tuple]) -> Optional[int]:
    for idx, row in enumerate(rows):
        texts = {_cell_text(c) for c in row}
        if HEADER_DATE in texts and HEADER_AMOUNT in texts:
            return idx
    return None


def parse_statement(data: bytes) -> dict:
    """Read a Max statement workbook. Raises ValueError when no sheet has the header row."""
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # not an xlsx at all
        raise ValueError("not a workbook") from exc

    holder: Optional[str] = None
    statement_month: Optional[str] = None
    card_last4: Optional[str] = None
    out_rows: list[dict] = []
    found_header = False

    try:
        for ws in wb.worksheets:
            rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
            header_idx = _find_header(rows)
            if header_idx is None:
                continue
            found_header = True

            # Title cells above the header: holder, then the statement month.
            for row in rows[:header_idx]:
                for cell in row:
                    text = _cell_text(cell)
                    if not text or not isinstance(cell, str):
                        continue
                    if holder is None:
                        holder = text.split("-", 1)[0].strip() or None
                    if statement_month is None and _MONTH_RE.match(text):
                        statement_month = text

            header = [_cell_text(c) for c in rows[header_idx]]
            col = {name: header.index(name) for name in header if name}
            date_col = col[HEADER_DATE]
            amount_col = col[HEADER_AMOUNT]

            def cell(row: tuple, key: str) -> Any:
                idx = col.get(HEADER_COLUMNS[key])
                return row[idx] if idx is not None and idx < len(row) else None

            for row in rows[header_idx + 1:]:
                raw_date = row[date_col] if date_col < len(row) else None
                if raw_date is None or _cell_text(raw_date) == "":
                    break
                iso = _parse_date(raw_date)
                amount = _parse_amount(row[amount_col] if amount_col < len(row) else None)
                if iso is None or amount is None:
                    continue
                if card_last4 is None:
                    last4 = _cell_text(cell(row, "card_last4"))
                    card_last4 = last4 or None
                out_rows.append({
                    "date": iso,
                    "merchant": normalise_merchant(cell(row, "merchant")),
                    "amount": amount,
                    "max_category": _cell_text(cell(row, "max_category")) or None,
                    "sheet": ws.title,
                })
    finally:
        wb.close()

    if not found_header:
        raise ValueError("no Max header row found")

    return {
        "holder": holder,
        "statement_month": statement_month,
        "card_last4": card_last4,
        "rows": out_rows,
    }


def classify(conn: sqlite3.Connection, rows: list[dict]) -> list[dict]:
    """Mark each file row `exists`, `recurring` or `new` (each DB row usable once).

    Exact date+amount matches are resolved first, in file order, so a
    ±3-day recurring match can never take a row another line matches exactly.
    """
    result = [dict(r, status="new", matched_id=None) for r in rows]
    if not result:
        return result

    dates = [date.fromisoformat(r["date"]) for r in result]
    lo = (min(dates) - timedelta(days=RECURRING_WINDOW_DAYS)).isoformat()
    hi = (max(dates) + timedelta(days=RECURRING_WINDOW_DAYS)).isoformat()
    existing = [
        {
            "id": r["id"],
            "date": r["date"],
            "abs": round(abs(float(r["amount"] or 0)), 2),
            "recurring": r["recurrence_id"] is not None,
        }
        for r in conn.execute(
            "SELECT id, date, amount, recurrence_id FROM transactions "
            "WHERE date >= ? AND date <= ? ORDER BY date, id",
            (lo, hi),
        ).fetchall()
    ]
    used: set[int] = set()

    # 1. exists — same date, same absolute amount.
    for row in result:
        amount = round(abs(row["amount"]), 2)
        match = next(
            (e for e in existing
             if e["id"] not in used and e["date"] == row["date"] and e["abs"] == amount),
            None,
        )
        if match:
            used.add(match["id"])
            row["status"] = "exists"
            row["matched_id"] = match["id"]

    # 2. recurring — a recurrence-materialised row, same amount, within ±3 days (closest wins).
    for row, day in zip(result, dates):
        if row["status"] != "new":
            continue
        amount = round(abs(row["amount"]), 2)
        candidates = []
        for e in existing:
            if e["id"] in used or not e["recurring"] or e["abs"] != amount:
                continue
            gap = abs((date.fromisoformat(e["date"]) - day).days)
            if gap <= RECURRING_WINDOW_DAYS:
                candidates.append((gap, e["id"]))
        if candidates:
            _gap, match_id = min(candidates)
            used.add(match_id)
            row["status"] = "recurring"
            row["matched_id"] = match_id

    return result


def guess_category(
    conn: sqlite3.Connection,
    merchant: str,
    max_category: Optional[str],
    categories: Iterable[Any],
) -> tuple[Optional[int], str]:
    """(category_id, source) with source one of merchant / map / fallback.

    `categories` are the categories the import may choose from (id, name, is_saving).
    """
    cats = [
        {"id": c["id"], "name": str(c["name"]).strip(), "is_saving": bool(c["is_saving"])}
        for c in categories
    ]
    by_id = {c["id"]: c for c in cats}
    by_name = {c["name"]: c for c in cats}

    if merchant:
        row = conn.execute(
            "SELECT category_id FROM transactions WHERE notes = ? "
            "ORDER BY date DESC, id DESC LIMIT 1",
            (merchant,),
        ).fetchone()
        if row and row["category_id"] in by_id:
            return row["category_id"], "merchant"

    for name in MAX_CATEGORY_MAP.get((max_category or "").strip(), ()):
        if name in by_name:
            return by_name[name]["id"], "map"

    if FALLBACK_CATEGORY in by_name:
        return by_name[FALLBACK_CATEGORY]["id"], "fallback"
    from ..api.transactions import INCOME_CATEGORIES  # lazy: the API imports this module

    plain = next(
        (c for c in cats if c["name"] not in INCOME_CATEGORIES and not c["is_saving"]),
        None,
    )
    return (plain["id"] if plain else None), "fallback"


def guess_payer(conn: sqlite3.Connection, holder: Optional[str], session_user: Any) -> Optional[int]:
    """The household member named first in `holder`, else the logged-in one, else the first."""
    people = household(conn)
    if not people:
        return None
    first_word = (holder or "").split()[0] if (holder or "").split() else ""
    if first_word:
        for person in people:
            if display_name(person["name"]) == first_word:
                return person["id"]
    me = find_person(people, session_user)
    if me:
        return me["id"]
    return people[0]["id"]
