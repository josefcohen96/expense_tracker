"""Card statement import (דוח האשראי): Max xlsx → preview → import → undo.

The workbooks are built here with openpyxl in the Max layout; the real
statement is never committed (it carries the holder's ID number).
Dates are in 2031 so nothing else in the test DB collides with them.
"""
import json
from io import BytesIO

from openpyxl import Workbook

MAX_HEADER = [
    "תאריך עסקה", "שם בית העסק", "קטגוריה", "4 ספרות אחרונות של כרטיס האשראי",
    "סוג עסקה", "סכום חיוב", "מטבע חיוב", "סכום עסקה מקורי", "מטבע עסקה מקורי",
    "תאריך חיוב", "הערות", "תיוגים", "מועדון הנחות", "מפתח דיסקונט",
    "אופן ביצוע ההעסקה", 'שער המרה ממטבע מקור/התחשבנות לש"ח',
]


def _max_workbook(rows_by_sheet, holder="יוסף כהן-000000000", month="03/2031", last4="4022"):
    """Write a Max-shaped statement: 3 title rows, header on row 4, data, blank, סך הכל, total.

    rows_by_sheet: {sheet title: [(date "DD-MM-YYYY", merchant, max_category, amount), ...]}
    """
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in rows_by_sheet.items():
        ws = wb.create_sheet(title)
        ws.append([holder])
        ws.append([f"{last4}-max בהצדעה"])
        ws.append([month])
        ws.append(MAX_HEADER)
        total = 0.0
        for day, merchant, category, amount in rows:
            total += amount
            ws.append([day, merchant, category, last4, "רגילה", amount, "₪", amount, "₪",
                       "02-04-2031", None, None, None, None, "אחר", None])
        ws.append([])
        ws.append(["סך הכל"])
        ws.append([f"{total:.2f}₪"])
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _preview(client, data):
    return client.post(
        "/api/transactions/import/preview",
        files={"file": ("transaction-details_export_1.xlsx", data,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def _category_id(db_conn, name):
    return db_conn.execute("SELECT id FROM categories WHERE name = ?", (name,)).fetchone()[0]


def _user_id(db_conn, name):
    return db_conn.execute("SELECT id FROM users WHERE name = ?", (name,)).fetchone()[0]


def _insert_tx(db_conn, day, amount, notes, category_id, recurrence_id=None, period_key=None):
    user_id = _user_id(db_conn, "Yosef")
    cur = db_conn.execute(
        "INSERT INTO transactions (date, amount, category_id, user_id, notes, recurrence_id, period_key) "
        "VALUES (?, ?, ?, ?, ?, ?, ?)",
        (day, amount, category_id, user_id, notes, recurrence_id, period_key),
    )
    db_conn.commit()
    return cur.lastrowid


def _delete_tx(db_conn, ids):
    for tx_id in ids:
        db_conn.execute("DELETE FROM transactions WHERE id = ?", (tx_id,))
    db_conn.commit()


def test_preview_parses_both_sheets(app_client):
    data = _max_workbook({
        "עסקאות במועד החיוב": [
            ("03-03-2031", "pytest-ci  WOLT   TLV", "מסעדות, קפה וברים", 205.9),
            ("04-03-2031", "pytest-ci SHUFERSAL", "מזון וצריכה", 34),
        ],
        'עסקאות חו"ל ומט"ח': [
            ("12-03-2031", "IHERB IHERB.COM        IHERB.COM     NL", "מזון וצריכה", 224.51),
        ],
    })
    r = _preview(app_client, data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["holder"] == "יוסף כהן"
    assert body["statement_month"] == "03/2031"
    assert body["card_last4"] == "4022"
    rows = body["rows"]
    assert [(x["date"], x["merchant"], x["amount"]) for x in rows] == [
        ("2031-03-03", "pytest-ci WOLT TLV", 205.9),
        ("2031-03-04", "pytest-ci SHUFERSAL", 34.0),
        ("2031-03-12", "IHERB IHERB.COM IHERB.COM NL", 224.51),
    ]
    assert [x["index"] for x in rows] == [0, 1, 2]
    assert rows[0]["max_category"] == "מסעדות, קפה וברים"
    assert all(x["amount"] > 0 for x in rows)
    assert body["counts"]["new"] + body["counts"]["exists"] + body["counts"]["recurring"] == 3


def test_preview_marks_existing_and_recurring(app_client, db_conn):
    cat = _category_id(db_conn, "פנאי")
    rec = app_client.post("/api/recurrences", json={
        "name": "pytest-ci-netflix",
        "amount": 49.9,
        "category_id": cat,
        "user_id": _user_id(db_conn, "Yosef"),
        "start_date": "2031-06-01",
        "frequency": "monthly",
        "day_of_month": 1,
        "active": True,
    })
    assert rec.status_code == 200, rec.text
    rec_id = rec.json()["id"]
    plain_id = _insert_tx(db_conn, "2031-05-10", -77.7, "pytest-ci-plain", cat)
    booked_id = _insert_tx(db_conn, "2031-05-01", -49.9, None, cat, rec_id, "2031-05")
    try:
        data = _max_workbook({"עסקאות במועד החיוב": [
            ("10-05-2031", "pytest-ci-plain", "פנאי, בידור וספורט", 77.7),     # same date + amount
            ("03-05-2031", "pytest-ci-netflix", "פנאי, בידור וספורט", 49.9),   # recurrence booked 2 days earlier
            ("11-05-2031", "pytest-ci-other", "פנאי, בידור וספורט", 77.7),     # plain row already used
            ("20-05-2031", "pytest-ci-late", "פנאי, בידור וספורט", 49.9),      # recurring row out of ±3 days
        ]})
        rows = _preview(app_client, data).json()["rows"]
        assert [(x["status"], x["matched_id"]) for x in rows] == [
            ("exists", plain_id),
            ("recurring", booked_id),
            ("new", None),
            ("new", None),
        ]
    finally:
        _delete_tx(db_conn, [plain_id, booked_id])
        app_client.delete(f"/api/recurrences/{rec_id}")


def test_preview_multiset_matching(app_client, db_conn):
    cat = _category_id(db_conn, "פנאי")
    tx_id = _insert_tx(db_conn, "2031-07-05", -30.0, "pytest-ci-coffee", cat)
    try:
        data = _max_workbook({"עסקאות במועד החיוב": [
            ("05-07-2031", "pytest-ci-coffee", "מסעדות, קפה וברים", 30),
            ("05-07-2031", "pytest-ci-coffee", "מסעדות, קפה וברים", 30),
        ]})
        body = _preview(app_client, data).json()
        assert [x["status"] for x in body["rows"]] == ["exists", "new"]
        assert body["counts"] == {"new": 1, "exists": 1, "recurring": 0}
    finally:
        _delete_tx(db_conn, [tx_id])


def test_preview_category_sources(app_client, db_conn):
    learned = _category_id(db_conn, "בריאות")
    older = _insert_tx(db_conn, "2030-01-01", -10.0, "pytest-ci WOLT", _category_id(db_conn, "פנאי"))
    newer = _insert_tx(db_conn, "2030-02-01", -10.0, "pytest-ci WOLT", learned)
    try:
        data = _max_workbook({"עסקאות במועד החיוב": [
            ("01-08-2031", "pytest-ci   WOLT", "מסעדות, קפה וברים", 11),
            ("02-08-2031", "pytest-ci-unknown-cafe", "מסעדות, קפה וברים", 12),
            ("03-08-2031", "pytest-ci-transfer", "העברת כספים", 13),
        ]})
        rows = _preview(app_client, data).json()["rows"]
        assert (rows[0]["category_source"], rows[0]["category_id"]) == ("merchant", learned)
        assert (rows[1]["category_source"], rows[1]["category_id"]) == ("map", _category_id(db_conn, "אוכל בחוץ"))
        assert (rows[2]["category_source"], rows[2]["category_id"]) == ("fallback", _category_id(db_conn, "הוצאות בית"))
    finally:
        _delete_tx(db_conn, [older, newer])


def test_preview_guesses_payer_from_holder(app_client, db_conn):
    rows = {"עסקאות במועד החיוב": [("01-09-2031", "pytest-ci-payer", "אופנה", 5)]}
    yosef = _preview(app_client, _max_workbook(rows, holder="יוסף כהן-000000000")).json()
    karina = _preview(app_client, _max_workbook(rows, holder="קארינה כהן-000000000")).json()
    stranger = _preview(app_client, _max_workbook(rows, holder="פלוני אלמוני-000000000"))
    assert yosef["user_id"] == _user_id(db_conn, "Yosef")
    assert karina["user_id"] == _user_id(db_conn, "Karina")
    assert karina["holder"] == "קארינה כהן"
    assert stranger.status_code == 200
    # Unknown holder → the logged-in user (Yosef under pytest).
    assert stranger.json()["user_id"] == _user_id(db_conn, "Yosef")
    account = db_conn.execute("SELECT id FROM accounts WHERE name = 'כרטיס אשראי'").fetchone()
    assert yosef["account_id"] == (account["id"] if account else None)


def test_import_then_repreview_is_idempotent(app_client, db_conn):
    data = _max_workbook({
        "עסקאות במועד החיוב": [
            ("05-10-2031", "pytest-ci  IMPORT-A", "מסעדות, קפה וברים", 205.9),
            ("06-10-2031", "pytest-ci IMPORT-B", "מזון וצריכה", -12.5),   # a refund
        ],
        'עסקאות חו"ל ומט"ח': [("07-10-2031", "pytest-ci IMPORT-C", "אופנה", 99)],
    })
    preview = _preview(app_client, data).json()
    assert preview["counts"]["new"] == 3
    karina = _user_id(db_conn, "Karina")
    account = db_conn.execute("SELECT id FROM accounts ORDER BY id LIMIT 1").fetchone()["id"]
    cat = _category_id(db_conn, "פנאי")
    payload = {
        "user_id": karina,
        "account_id": account,
        "rows": [{"date": x["date"], "amount": x["amount"], "merchant": x["merchant"],
                  "category_id": cat if i == 2 else x["category_id"]}
                 for i, x in enumerate(preview["rows"])],
    }
    r = app_client.post("/api/transactions/import", json=payload)
    assert r.status_code == 200, r.text
    body = r.json()
    ids = body["created"]
    try:
        assert body["count"] == 3 and len(ids) == 3
        assert (body["date_from"], body["date_to"]) == ("2031-10-05", "2031-10-07")
        placeholders = ",".join("?" * len(ids))
        saved = db_conn.execute(
            f"SELECT * FROM transactions WHERE id IN ({placeholders}) ORDER BY id", ids
        ).fetchall()
        assert [(s["date"], s["amount"], s["notes"]) for s in saved] == [
            ("2031-10-05", -205.9, "pytest-ci IMPORT-A"),
            ("2031-10-06", 12.5, "pytest-ci IMPORT-B"),
            ("2031-10-07", -99.0, "pytest-ci IMPORT-C"),
        ]
        assert all(s["user_id"] == karina and s["account_id"] == account for s in saved)
        assert all(s["recurrence_id"] is None and s["tags"] is None for s in saved)
        assert saved[2]["category_id"] == cat
        assert saved[0]["category_id"] == preview["rows"][0]["category_id"]

        again = _preview(app_client, data).json()
        assert [x["status"] for x in again["rows"]] == ["exists", "exists", "exists"]
        assert again["counts"] == {"new": 0, "exists": 3, "recurring": 0}
        # The merchant memory now answers for C.
        assert (again["rows"][2]["category_source"], again["rows"][2]["category_id"]) == ("merchant", cat)
    finally:
        _delete_tx(db_conn, ids)


def test_import_rejects_bad_payload(app_client, db_conn):
    cat = _category_id(db_conn, "פנאי")
    yosef = _user_id(db_conn, "Yosef")
    row = {"date": "2031-10-10", "amount": 5, "merchant": "pytest-ci-bad", "category_id": cat}
    assert app_client.post("/api/transactions/import", json={"user_id": yosef, "rows": []}).status_code == 400
    assert app_client.post("/api/transactions/import",
                           json={"user_id": 999999, "rows": [row]}).status_code == 400
    assert app_client.post("/api/transactions/import",
                           json={"user_id": yosef, "rows": [dict(row, category_id=999999)]}).status_code == 400
    left = db_conn.execute("SELECT COUNT(*) FROM transactions WHERE notes = 'pytest-ci-bad'").fetchone()[0]
    assert left == 0


def test_undo_deletes_only_imported(app_client, db_conn):
    cat = _category_id(db_conn, "פנאי")
    yosef = _user_id(db_conn, "Yosef")
    rec = app_client.post("/api/recurrences", json={
        "name": "pytest-ci-undo-rec", "amount": 20.0, "category_id": cat, "user_id": yosef,
        "start_date": "2031-12-01", "frequency": "monthly", "day_of_month": 1, "active": True,
    })
    assert rec.status_code == 200, rec.text
    rec_id = rec.json()["id"]
    booked = _insert_tx(db_conn, "2031-11-01", -20.0, None, cat, rec_id, "2031-11")
    bystander = _insert_tx(db_conn, "2031-11-02", -8.0, "pytest-ci-bystander", cat)
    try:
        r = app_client.post("/api/transactions/import", json={
            "user_id": yosef, "account_id": None,
            "rows": [
                {"date": "2031-11-03", "amount": 15, "merchant": "pytest-ci-undo-1", "category_id": cat},
                {"date": "2031-11-04", "amount": 16, "merchant": "pytest-ci-undo-2", "category_id": cat},
            ],
        })
        assert r.status_code == 200, r.text
        ids = r.json()["created"]
        u = app_client.post("/api/transactions/import/undo", json={"ids": ids + [booked]})
        assert u.status_code == 200, u.text
        assert u.json() == {"deleted": 2}
        exists = lambda i: db_conn.execute("SELECT 1 FROM transactions WHERE id = ?", (i,)).fetchone() is not None
        assert not any(exists(i) for i in ids)
        assert exists(booked)
        assert exists(bystander)
    finally:
        _delete_tx(db_conn, [booked, bystander])
        app_client.delete(f"/api/recurrences/{rec_id}")


def test_preview_rejects_unknown_file(app_client):
    wb = Workbook()
    wb.active.append(["שם", "סכום"])
    wb.active.append(["משהו", 10])
    bio = BytesIO()
    wb.save(bio)
    r = _preview(app_client, bio.getvalue())
    assert r.status_code == 400
    assert r.json()["detail"] == "הקובץ לא נראה כמו דוח של max"

    not_xlsx = _preview(app_client, b"just some text")
    assert not_xlsx.status_code == 400
    assert not_xlsx.json()["detail"] == "הקובץ לא נראה כמו דוח של max"


def test_import_page_and_link(app_client):
    page = app_client.get("/finances/transactions/import")
    assert page.status_code == 200
    assert "דוח האשראי" in page.text
    assert "דוח מקס (xlsx) — גרור לכאן או לחץ לבחירה" in page.text
    # Expense categories only, as on the transactions page (the list is embedded as JSON).
    assert json.dumps("אוכל בחוץ") in page.text
    assert json.dumps("משכורת") not in page.text and "משכורת" not in page.text

    listing = app_client.get("/finances/transactions")
    assert listing.status_code == 200
    assert 'href="/finances/transactions/import"' in listing.text
    assert "טען דוח אשראי" in listing.text


# ─── Part 2: the mailbox import (POST /api/transactions/import/auto) ──────────

def _auto(client, data, token=None):
    headers = {"Authorization": f"Bearer {token}"} if token is not None else {}
    return client.post(
        "/api/transactions/import/auto",
        files={"file": ("transaction-details_export_1.xlsx", data,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )


def _count_notes(db_conn, prefix):
    return db_conn.execute(
        "SELECT COUNT(*) FROM transactions WHERE notes LIKE ?", (prefix + "%",)
    ).fetchone()[0]


def test_auto_import_disabled_without_token(app_client, db_conn, monkeypatch):
    monkeypatch.delenv("IMPORT_TOKEN", raising=False)
    data = _max_workbook({"עסקאות במועד החיוב": [
        ("03-01-2032", "pytest-ci-auto-off", "מזון וצריכה", 10),
    ]})
    r = _auto(app_client, data, token="anything")
    assert r.status_code == 403
    assert r.json()["detail"] == "ייבוא אוטומטי כבוי"
    monkeypatch.setenv("IMPORT_TOKEN", "")
    r = _auto(app_client, data, token="")
    assert r.status_code == 403
    assert r.json()["detail"] == "ייבוא אוטומטי כבוי"
    assert _count_notes(db_conn, "pytest-ci-auto-off") == 0


def test_auto_import_rejects_bad_token(app_client, db_conn, monkeypatch):
    monkeypatch.setenv("IMPORT_TOKEN", "pytest-ci-secret-token")
    data = _max_workbook({"עסקאות במועד החיוב": [
        ("04-01-2032", "pytest-ci-auto-bad", "מזון וצריכה", 10),
    ]})
    missing = _auto(app_client, data)
    wrong = _auto(app_client, data, token="pytest-ci-wrong-token")
    not_bearer = app_client.post(
        "/api/transactions/import/auto",
        files={"file": ("x.xlsx", data, "application/octet-stream")},
        headers={"Authorization": "Basic pytest-ci-secret-token"},
    )
    for r in (missing, wrong, not_bearer):
        assert r.status_code == 401
        assert r.json() == {"detail": "אסימון לא תקין"}
    assert _count_notes(db_conn, "pytest-ci-auto-bad") == 0


def test_auto_import_adds_only_new_and_is_idempotent(app_client, db_conn, monkeypatch):
    monkeypatch.setenv("IMPORT_TOKEN", "pytest-ci-secret-token")
    cat = _category_id(db_conn, "פנאי")
    existing = _insert_tx(db_conn, "2032-02-05", -40.0, "pytest-ci-auto-known", cat)
    data = _max_workbook({
        "עסקאות במועד החיוב": [
            ("05-02-2032", "pytest-ci-auto-known", "פנאי, בידור וספורט", 40),   # already in the app
            ("06-02-2032", "pytest-ci-auto  CAFE", "מסעדות, קפה וברים", 55.5),
            ("07-02-2032", "pytest-ci-auto-transfer", "העברת כספים", 20),       # no category → fallback
        ],
        'עסקאות חו"ל ומט"ח': [("09-02-2032", "pytest-ci-auto-refund", "אופנה", -15)],
    }, holder="קארינה כהן-000000000", month="02/2032")
    ids = []
    try:
        r = _auto(app_client, data, token="pytest-ci-secret-token")
        assert r.status_code == 200, r.text
        body = r.json()
        ids = body["created"]
        assert body["added"] == 3 and len(ids) == 3
        assert body["skipped"] == {"exists": 1, "recurring": 0}
        assert body["unsorted"] == 1
        assert body["holder"] == "קארינה כהן"
        assert body["statement_month"] == "02/2032"
        assert (body["date_from"], body["date_to"]) == ("2032-02-06", "2032-02-09")

        placeholders = ",".join("?" * len(ids))
        saved = db_conn.execute(
            f"SELECT * FROM transactions WHERE id IN ({placeholders}) ORDER BY id", ids
        ).fetchall()
        assert [(s["date"], s["amount"], s["notes"]) for s in saved] == [
            ("2032-02-06", -55.5, "pytest-ci-auto CAFE"),
            ("2032-02-07", -20.0, "pytest-ci-auto-transfer"),
            ("2032-02-09", 15.0, "pytest-ci-auto-refund"),
        ]
        karina = _user_id(db_conn, "Karina")
        account = db_conn.execute("SELECT id FROM accounts WHERE name = 'כרטיס אשראי'").fetchone()
        assert all(s["user_id"] == karina for s in saved)
        assert all(s["account_id"] == (account["id"] if account else None) for s in saved)
        assert all(s["recurrence_id"] is None and s["tags"] is None for s in saved)
        assert saved[0]["category_id"] == _category_id(db_conn, "אוכל בחוץ")
        assert saved[1]["category_id"] == _category_id(db_conn, "הוצאות בית")
        assert _count_notes(db_conn, "pytest-ci-auto-known") == 1

        again = _auto(app_client, data, token="pytest-ci-secret-token")
        assert again.status_code == 200, again.text
        second = again.json()
        assert second["added"] == 0 and second["created"] == []
        assert second["skipped"] == {"exists": 4, "recurring": 0}
        assert second["unsorted"] == 0
        assert (second["date_from"], second["date_to"]) == (None, None)
        assert _count_notes(db_conn, "pytest-ci-auto") == 4
    finally:
        _delete_tx(db_conn, ids + [existing])


def test_auto_import_route_is_public():
    from app.backend.app.auth import build_public_route_matchers
    from app.backend.app.main import app

    matchers = build_public_route_matchers(app)
    assert any(
        regex.match("/api/transactions/import/auto") and "POST" in methods
        for regex, methods in matchers
    )
    # The rest of the import API still needs a session.
    assert not any(
        regex.match("/api/transactions/import") or regex.match("/api/transactions/import/preview")
        for regex, _methods in matchers
    )
